"""Xuất bảng theo mẫu "Phụ lục 6.4" như sổ Excel đang dùng (mục 5.7 HUONG_DAN.md)."""
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

from .phan_tich import SO_LA_MA

TIEU_DE_MAC_DINH = "Phụ lục 6.4. Tổng hợp giấy phép môi trường do UBND tỉnh cấp phép, hoặc UBND tỉnh ủy quyền cấp phép"
COT = ["TT", "Tên cơ sở", "Chủ cơ sở", "Địa chỉ", "Loại hình sản xuất", "Giấy phép số", "Công suất",
       "Nước thải (m3/ngày đêm)", "Khí thải (m3/giờ)", "Chất thải rắn thông thường (kg/năm)",
       "Chất thải nguy hại (kg/năm)", "Trạng thái"]
RONG = [6, 38, 30, 34, 22, 26, 28, 16, 14, 18, 16, 18]


def _gia_tri(ds, loai):
    """Ưu tiên nguyên văn để khớp với sổ; nhiều dòng (vd 2 dòng khí thải) nối bằng '; '."""
    return "; ".join(x.luu_luong_goc if hasattr(x, "luu_luong_goc") else x.khoi_luong_goc
                     for x in ds if x.loai == loai and (getattr(x, "luu_luong_goc", None) or
                                                         getattr(x, "khoi_luong_goc", None)))


def _gp_so(gp):
    if not gp.so_hieu:
        return gp.so_hieu_goc or ""
    return f"{gp.so_hieu} ngày {gp.ngay_ky.strftime('%d/%m/%Y')}" if gp.ngay_ky else gp.so_hieu


def xuat_phu_luc_64(ds_gpmt, tep, tieu_de=None, noi_bo=True):
    wb = Workbook()
    ws = wb.active
    ws.title = "Phu luc 6.4"
    manh = Side(style="thin")
    khung = Border(left=manh, right=manh, top=manh, bottom=manh)
    dam = Font(bold=True)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(COT))
    ws.cell(1, 1, tieu_de or TIEU_DE_MAC_DINH).font = Font(bold=True, size=13)
    ws.cell(1, 1).alignment = Alignment(horizontal="center", wrap_text=True)
    ws.row_dimensions[1].height = 34
    for i, ten in enumerate(COT, 1):
        o = ws.cell(3, i, ten)
        o.font, o.border = dam, khung
        o.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = RONG[i - 1]
    ws.row_dimensions[3].height = 45

    theo_nam = defaultdict(list)
    for gp in ds_gpmt:
        theo_nam[gp.nam_cap].append(gp)
    r = 4
    for stt_nhom, nam in enumerate(sorted(theo_nam, key=lambda x: (x is None, x or 0)), 1):
        la_ma = SO_LA_MA[stt_nhom] if stt_nhom < len(SO_LA_MA) else str(stt_nhom)
        ws.cell(r, 1, la_ma).font = dam
        ws.cell(r, 2, f"Năm {nam}" if nam else "Chưa xác định năm").font = dam
        for c in range(1, len(COT) + 1):
            ws.cell(r, c).border = khung
        r += 1
        for tt, gp in enumerate(sorted(theo_nam[nam], key=lambda g: (g.ngay_ky is None, g.ngay_ky or 0)), 1):
            cs = gp.co_so
            dong = [tt, cs.ten if cs else "", str(cs.chu_the) if cs and cs.chu_the else "",
                    (cs.xa_phuong_moi or cs.dia_diem or "") if cs else "", cs.loai_hinh if cs else "",
                    _gp_so(gp), gp.cong_suat_goc or (cs.cong_suat_goc if cs else "") or "",
                    _gia_tri(gp.xa_thai, "nước thải"), _gia_tri(gp.xa_thai, "khí thải"),
                    _gia_tri(gp.chat_thai, "CTR thông thường (sổ)") or "; ".join(
                        x for x in (_gia_tri(gp.chat_thai, "CTRCNTT"), _gia_tri(gp.chat_thai, "CTRSH")) if x),
                    _gia_tri(gp.chat_thai, "CTNH"), gp.trang_thai]
            for c, v in enumerate(dong, 1):
                o = ws.cell(r, c, v)
                o.border = khung
                o.alignment = Alignment(vertical="top", wrap_text=True)
            r += 1
    ws.freeze_panes = "A4"
    if noi_bo:
        ws2 = wb.create_sheet("Can ra soat")
        ws2.append(["Số hiệu", "Cơ sở", "Nguồn dữ liệu", "Lý do cần rà soát"])
        for gp in ds_gpmt:
            if gp.can_ra_soat:
                ws2.append([gp.so_hieu or "(chưa có số)", gp.co_so.ten if gp.co_so else "",
                            gp.nguon_du_lieu, "\n".join(gp.ds_ly_do)])
        for i, w in enumerate([20, 50, 40, 80], 1):
            ws2.column_dimensions[get_column_letter(i)].width = w
    wb.save(tep)
