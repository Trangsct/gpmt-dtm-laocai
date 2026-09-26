"""Chế độ công khai (Bạn chốt 27/9/2026) và trang tạo tài khoản quản trị đầu tiên."""
import io
import re

from openpyxl import load_workbook

from gpmt.models import Gpmt, TaiKhoan

from .conftest import tao_tai_khoan


def test_khach_xem_duoc_khong_can_dang_nhap(app, db, du_lieu_that):
    tao_tai_khoan(db, "qt@thu.local", "quan_tri")
    c = app.test_client()
    gp = db.session.query(Gpmt).filter_by(so_hieu="1439/GPMT-UBND").one()
    for u in ("/", "/thong-ke", "/gpmt/", "/gpmt/?q=Eurostark", f"/gpmt/{gp.id}", f"/co-so/{gp.co_so_du_an_id}", "/co-so",
              "/dtm", "/dang-ky-mt", "/chu-the"):
        assert c.get(u).status_code == 200, u
    r = c.get(f"/gpmt/{gp.id}")
    assert "Hết hiệu lực — bị thay thế" in r.text
    assert "Cán bộ đăng nhập" in r.text


def test_khach_khong_thay_noi_bo_va_khong_sua(app, db, du_lieu_that):
    tao_tai_khoan(db, "qt@thu.local", "quan_tri")
    c = app.test_client()
    co_co = db.session.query(Gpmt).filter(Gpmt.can_ra_soat.is_(True)).first()
    r = c.get(f"/gpmt/{co_co.id}")
    assert "Cần rà soát" not in r.text and "Nguồn dữ liệu" not in r.text and ">Sửa<" not in r.text
    for u in ("/ho-so", "/ra-soat", "/nhat-ky", "/nhap-du-lieu/", "/tai-khoan/", f"/gpmt/{co_co.id}/sua", "/gpmt/moi"):
        r = c.get(u)
        assert r.status_code == 302 and "/dang-nhap" in r.headers["Location"], u
    assert "Hồ sơ đang giải quyết" not in c.get("/").text
    assert "Hồ sơ đang giải quyết" not in c.get("/thong-ke").text


def test_khach_xuat_excel_khong_co_sheet_ra_soat(app, db, du_lieu_that):
    tao_tai_khoan(db, "qt@thu.local", "quan_tri")
    wb = load_workbook(io.BytesIO(app.test_client().get("/xuat-excel").data))
    assert wb.sheetnames == ["Phu luc 6.4"]


def test_tao_quan_tri_dau_tien_tren_web(app, db):
    c = app.test_client()
    assert c.get("/dang-nhap").headers["Location"].endswith("/cai-dat")
    tok = re.search(r'name="_csrf" value="([^"]+)"', c.get("/cai-dat").text).group(1)
    c.post("/cai-dat", data={"_csrf": tok, "ho_ten": "A", "email": "a@so.local", "mat_khau": "ngan",
                             "mat_khau_2": "ngan"})
    assert db.session.query(TaiKhoan).count() == 0
    r = c.post("/cai-dat", data={"_csrf": tok, "ho_ten": "Trần Văn A", "email": "A@So.local",
                                 "mat_khau": "mat-khau-dai-123", "mat_khau_2": "mat-khau-dai-123"})
    assert r.status_code == 302 and "/nhap-du-lieu/" in r.headers["Location"]
    tk = db.session.query(TaiKhoan).one()
    assert tk.email == "a@so.local" and tk.vai_tro == "quan_tri"
    assert c.get("/nhap-du-lieu/").status_code == 200          # đã đăng nhập luôn
    # Có tài khoản rồi: trang cài đặt khóa, không tạo thêm được
    from flask import g
    g.pop("_login_user", None)   # fixture giữ chung app context: bỏ người dùng của trình duyệt trước
    c2 = app.test_client()
    assert c2.get("/cai-dat").headers["Location"].endswith("/dang-nhap")
    tok2 = re.search(r'name="_csrf" value="([^"]+)"', c2.get("/dang-nhap").text).group(1)
    c2.post("/cai-dat", data={"_csrf": tok2, "ho_ten": "X", "email": "x@x.x", "mat_khau": "mat-khau-dai-123",
                              "mat_khau_2": "mat-khau-dai-123"})
    assert db.session.query(TaiKhoan).count() == 1


def test_trang_chu_gioi_thieu(app, db, du_lieu_that):
    r = app.test_client().get("/")
    assert r.status_code == 200
    for x in ("Sở Nông nghiệp và Môi trường", "Tra cứu", "Thủ tục hành chính", "736/QĐ-UBND", "22 ngày",
              "2519/GPMT-UBND", "số 64 đường Lý Tự Trọng", "02143.820 062",
              "contact-snnmt@laocai.gov.vn", "Trần Minh Sáng", "snnmt.laocai.gov.vn", "data-co=\"1.4\""):
        assert x in r.text, x
