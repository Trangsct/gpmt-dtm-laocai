"""Đăng nhập, phân quyền, nhật ký, xuất Excel."""
import io
import json

from openpyxl import load_workbook

from gpmt.models import Gpmt, NhatKy

from .conftest import dang_nhap, tao_tai_khoan


def test_phai_dang_nhap(app):
    assert app.test_client().get("/gpmt/").status_code == 302


def test_csrf_chan_post_thieu_token(app, db):
    tao_tai_khoan(db, "bt@thu.local", "bien_tap")
    c = app.test_client()
    dang_nhap(c, "bt@thu.local")
    assert c.post("/co-so/moi", data={"ten": "X"}).status_code == 400


def test_sua_ghi_nhat_ky(app, db, du_lieu_that):
    tao_tai_khoan(db, "bt@thu.local", "bien_tap")
    c = app.test_client()
    tok = dang_nhap(c, "bt@thu.local")
    gp = db.session.query(Gpmt).filter_by(so_hieu="2074/GPMT-UBND").one()
    r = c.get(f"/gpmt/{gp.id}/sua")
    assert r.status_code == 200
    du_lieu = {"_csrf": tok, "so_hieu": "2074/GPMT-UBND", "ngay_ky": "2022-09-20", "co_quan_cap": "UBND tỉnh Lào Cai",
               "loai_van_ban": "GPMT", "co_so_du_an_id": str(gp.co_so_du_an_id), "thoi_han_nam": "7",
               "cong_khai": "1"}
    r = c.post(f"/gpmt/{gp.id}/sua", data=du_lieu)
    assert r.status_code == 302
    db.session.expire_all()
    gp = db.session.get(Gpmt, gp.id)
    assert gp.thoi_han_nam == 7 and gp.ngay_het_han.year == 2029
    nk = db.session.query(NhatKy).filter_by(bang="gpmt", ban_ghi=gp.id, hanh_dong="sua").one()
    assert json.loads(nk.thay_doi)["thoi_han_nam"] == [None, 7.0]


def test_chi_xem_khong_duoc_sua(app, db, du_lieu_that):
    tao_tai_khoan(db, "xem@thu.local", "chi_xem")
    c = app.test_client()
    dang_nhap(c, "xem@thu.local")
    gp = db.session.query(Gpmt).first()
    assert c.get(f"/gpmt/{gp.id}").status_code == 200
    assert c.get(f"/gpmt/{gp.id}/sua").status_code == 403


def test_tai_khoan_xa_chi_thay_dia_ban_minh(app, db, du_lieu_that):
    tao_tai_khoan(db, "xa@thu.local", "ben_ngoai", pham_vi_xa="Văn Phú")
    c = app.test_client()
    dang_nhap(c, "xa@thu.local")
    r = c.get("/gpmt/")
    assert "2519/GPMT-UBND" in r.text and "2074/GPMT-UBND" not in r.text
    ngoai = db.session.query(Gpmt).filter_by(so_hieu="2074/GPMT-UBND").one()
    assert c.get(f"/gpmt/{ngoai.id}").status_code == 403
    for u in ("/ho-so", "/ra-soat", "/nhat-ky", "/tai-khoan/"):
        assert c.get(u).status_code == 403


def test_xuat_phu_luc_64(app, db, du_lieu_that):
    tao_tai_khoan(db, "xem@thu.local", "chi_xem")
    c = app.test_client()
    dang_nhap(c, "xem@thu.local")
    r = c.get("/xuat-excel?dia_ban=YB")
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.data))
    ws = wb["Phu luc 6.4"]
    assert ws.cell(1, 1).value.startswith("Phụ lục 6.4")
    gia_tri = [ws.cell(r, 6).value for r in range(4, ws.max_row + 1)]
    assert "1439/GPMT-UBND ngày 22/08/2022" in gia_tri
    assert "Can ra soat" in wb.sheetnames


def test_mo_pdf_goc(app, db, du_lieu_that):
    tao_tai_khoan(db, "xem@thu.local", "chi_xem")
    c = app.test_client()
    dang_nhap(c, "xem@thu.local")
    gp = db.session.query(Gpmt).filter_by(so_hieu="2519/GPMT-UBND").one()
    r = c.get(f"/gpmt/{gp.id}/pdf")
    assert r.status_code == 200 and r.data[:4] == b"%PDF"
